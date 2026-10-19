"""Build consistent PPTX/PDF slides and a ZIP from MATLAB's measured outputs.

Without summary.json this creates a visibly marked draft, never invented results.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
from pathlib import Path
import zipfile

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parent
SIZE = (1600, 900)
BG = "#101c30"
FG = "#edf3fa"
MUTED = "#a9bed6"
ACCENT = "#52d6c7"
PANEL = "#1c304b"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    directory = Path(r"C:\Windows\Fonts")
    candidates = [
        directory / ("arialbd.ttf" if bold else "arial.ttf"),
        Path("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"),
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(str(path), size)
        except OSError:
            continue
    raise RuntimeError("Install Arial or DejaVu Sans to render slides.")


def rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def flag(value: str) -> bool:
    return value.strip().lower() in {"1", "true"}


def wrapped(draw: ImageDraw.ImageDraw, text: str, xy: tuple[int, int],
            width: int, size: int = 28, fill: str = FG, bold: bool = False) -> int:
    x, y = xy
    face = font(size, bold)
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if line and draw.textlength(candidate, font=face) > width:
                draw.text((x, y), line, font=face, fill=fill)
                y += round(size * 1.35)
                line = word
            else:
                line = candidate
        draw.text((x, y), line, font=face, fill=fill)
        y += round(size * 1.5)
    return y


class Deck:
    def __init__(self, draft: bool):
        self.draft = draft
        self.pages: list[Image.Image] = []
        self.notes: list[str] = []

    def slide(self, title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
        image = Image.new("RGB", SIZE, BG)
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, 1599, 12), fill=ACCENT)
        wrapped(draw, title, (65, 50), 1470, 46, bold=True)
        if subtitle:
            wrapped(draw, subtitle, (68, 122), 1460, 24, MUTED)
        status = "DRAFT - awaiting reviewed labels and MATLAB results" if self.draft else \
            "Measured experiment | replace contribution placeholders before submission"
        draw.text((65, 853), status, font=font(19), fill=MUTED)
        draw.text((1500, 853), str(len(self.pages) + 1), font=font(19), fill=MUTED)
        self.pages.append(image)
        self.notes.append("")
        return image, draw

    def note(self, text: str) -> None:
        self.notes[-1] = text

    def card(self, draw: ImageDraw.ImageDraw, x: int, y: int, w: int, h: int,
             title: str, text: str, size: int = 27) -> None:
        draw.rounded_rectangle((x, y, x + w, y + h), 18, fill=PANEL)
        bottom = wrapped(draw, title, (x + 25, y + 22), w - 50, 30, ACCENT, True)
        bottom = wrapped(draw, text, (x + 25, bottom + 12), w - 50, size)
        if bottom > y + h - 5:
            raise ValueError(f"Text overflows slide card: {title}")


def paste_image(canvas: Image.Image, path: Path, box: tuple[int, int, int, int]) -> None:
    with Image.open(path) as image:
        fitted = ImageOps.contain(image.convert("RGB"), (box[2], box[3]))
        canvas.paste(fitted, (box[0] + (box[2] - fitted.width) // 2,
                             box[1] + (box[3] - fitted.height) // 2))


def percent(value: float) -> str:
    return f"{100 * value:.1f}%"


def metric_text(metrics: dict) -> str:
    p = percent(metrics["Precision"]) if metrics["PrecisionDefined"] else "undefined"
    r = percent(metrics["Recall"]) if metrics["RecallDefined"] else "undefined"
    t = metrics["Threshold"]
    threshold = "reject-all (Inf)" if t is None else f"{t:.4g}"
    return (
        f"Precision: {p} | Recall: {r}\n"
        f"TP {metrics['TP']} | FP {metrics['FP']} | FN {metrics['FN']}\n"
        f"Localization: {metrics['CorrectLocalization']}/{metrics['AcceptedDetections']} "
        f"({percent(metrics['LocalizationSuccess']) if metrics['LocalizationSuccessDefined'] else 'undefined'})\n"
        f"Threshold: {threshold}\n"
        f"Runtime: {metrics['MeanSecondsPerFrame']:.3f} s/frame\n"
        f"Targets met on evaluation: {metrics['EvaluationTargetsMet']}"
    )


def evidence(deck: Deck, root: Path, labels: list[dict], results: Path,
             categories: list[tuple[str, str]], title: str) -> None:
    image, draw = deck.slide(title, "Actual held-out predictions only; no synthetic success/failure examples.")
    outcomes = rows(results / "final_outcomes.csv")
    label_map = {row["Frame"]: row for row in labels}
    label_index = {row["Frame"]: k + 1 for k, row in enumerate(labels)}
    for column, (key, caption) in enumerate(categories):
        x = 65 + column * 755
        selected = None
        for row in outcomes:
            if key == "replica":
                reference = label_map[row["Frame"]]
                matches = flag(row["TN"]) and reference["Distractor"].lower() == "printed replica"
            else:
                matches = flag(row[key])
            if matches:
                selected = row
                break
        draw.rounded_rectangle((x, 195, x + 720, 750), 18, fill=PANEL)
        wrapped(draw, caption, (x + 20, 213), 680, 29, ACCENT, True)
        if selected:
            k = label_index[selected["Frame"]]
            path = results / "visualizations" / f"frame_{k:03d}.jpg"
            if path.is_file():
                paste_image(image, path, (x + 15, 280, 690, 390))
                wrapped(draw, Path(selected["Frame"].replace("\\", "/")).name,
                        (x + 20, 687), 680, 18, MUTED)
            else:
                wrapped(draw, "Measured example found. Run exportDemo to render its overlay.",
                        (x + 30, 350), 650, 30)
        else:
            text = "Pending measured evaluation." if deck.draft else \
                "No example of this outcome in the evaluated frames. Do not manufacture one."
            wrapped(draw, text, (x + 30, 350), 650, 30)
    deck.note("Cyan is the human reference; yellow is the baseline; green is the final acceptance. "
              "A misplaced acceptance can be both FP and FN. Discuss only examples actually observed.")


def threshold_plot(image: Image.Image, draw: ImageDraw.ImageDraw, results: Path) -> bool:
    data = rows(results / "final_thresholds.csv")
    finite = [r for r in data if math.isfinite(float(r["Threshold"]))]
    if not finite:
        return False
    x0, y0, width, height = 115, 680, 810, 370
    maximum = max(float(r["Threshold"]) for r in finite) or 1
    draw.line((x0, y0 - height, x0, y0, x0 + width, y0), fill=MUTED, width=3)
    for field, color in [("Precision", ACCENT), ("Recall", "#ffb75d")]:
        points = [(x0 + int(width * float(r["Threshold"]) / maximum),
                   y0 - int(height * float(r[field]))) for r in finite]
        if len(points) > 1:
            draw.line(points, fill=color, width=4)
        else:
            x, y = points[0]
            draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill=color)
    draw.text((x0, y0 + 20), "0", fill=MUTED, font=font(23))
    draw.text((x0 + width - 90, y0 + 20), f"{maximum:.2g}", fill=MUTED, font=font(23))
    draw.text((x0, y0 - height - 45), "1.0     Precision (teal) / Recall (orange)",
              fill=MUTED, font=font(24))
    draw.text((x0 + 250, y0 + 65), "SIFT acceptance threshold", fill=MUTED, font=font(24))
    return True


def build_deck(root: Path, results: Path) -> Deck:
    summary_path = results / "summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8")) if summary_path.exists() else None
    if summary and summary.get("Status") != "Measured on reviewed held-out frames":
        raise ValueError("Only a measured MATLAB summary can populate the results slides.")
    draft = summary is None
    labels = rows(results / "labels_snapshot.csv")
    deck = Deck(draft)
    image, draw = deck.slide("Anchor-Based Recognition & Localization",
                            "Project 3 | instance recognition, not category detection")
    deck.card(draw, 65, 220, 710, 470, "The question",
              "Does this frame contain the SAME physical extinguisher shown in three anchors?\n"
              "Return one ranked, accepted region or target-absent.\n"
              "A printed image of the target is a negative.", 32)
    anchor = root / "Anchor_Images" / "Anchor1.png"
    if anchor.exists():
        paste_image(image, anchor, (850, 185, 620, 610))
    deck.note("Introduce the physical instance and printed-replica challenge. "
              "The anchors define the target; there is no trained category detector. "
              "This is a draft unless results have been generated.")

    _, draw = deck.slide("Design: one candidate generator, two recognition methods")
    deck.card(draw, 65, 210, 710, 500, "Simple baseline",
              "Red connected components propose regions.\n"
              "Expand each region to include the hose.\n"
              "Compare normalized 16 x 4 HSV histograms with all anchors.\n"
              "Rank by Bhattacharyya coefficient; accept the best above a calibrated threshold.")
    deck.card(draw, 815, 210, 720, 500, "Final system",
              "Use the SAME proposal regions.\n"
              "SIFT descriptors + unique ratio-test matching.\n"
              "RANSAC homography + geometry/support checks.\n"
              "Rank valid region-anchor pairs; accept best score above a separate threshold.")
    deck.note("A shared proposal generator makes the matching comparison clearer, but proposal misses "
              "affect both. Histogram similarity alone ignores spatial arrangement and physical identity.")

    _, draw = deck.slide("System architecture")
    for i, (heading, text) in enumerate([
        ("Anchors", "Human object crops\nResize + SIFT\nHSV descriptors"),
        ("Frame", "Resize RGB\nRed-mask proposals\nExpand to hose"),
        ("Compare", "HSV or SIFT\nGeometric checks\nRank candidates"),
        ("Output", "Frozen threshold\nPresent / absent\nRegion + evidence"),
    ]):
        x = 65 + i * 380
        deck.card(draw, x, 260, 340, 330, heading, text, 28)
        if i < 3:
            draw.line((x + 345, 420, x + 375, 420), fill=ACCENT, width=5)
    wrapped(draw, "Validation labels -> threshold calibration -> freeze -> held-out evaluation",
            (100, 685), 1400, 32, ACCENT)
    deck.note("Explain every arrow. Ground truth never enters feature extraction or matching. "
              "Only validation frames choose the thresholds. Evaluation reports the frozen system.")

    _, draw = deck.slide("Matching evidence and localization")
    deck.card(draw, 65, 205, 710, 555, "SIFT evidence",
              "Match ratio <= 0.75; unique correspondences.\n"
              "Projective RANSAC: 2,000 trials; 4 px residual limit in resized candidate.\n"
              "At least 6 inliers; inlier ratio >= 0.35.\n"
              "Anchor inlier bounding-span coverage >= 0.02.\n"
              "Score = inliers x inlier ratio x sqrt(coverage).", 27)
    deck.card(draw, 815, 205, 720, 555, "Reject impossible geometry",
              "Finite convex projected quadrilateral.\n"
              "Area, size, frame and proposal-center checks.\n"
              "Transform cropped anchor corners to the original frame.\n"
              "Output enclosing axis-aligned box and center.\n"
              "No valid proposal is absent even at threshold zero.", 27)
    deck.note("Coverage is a bounding-span fraction, not a convex-hull area. Homography is an approximate "
              "model for this nonplanar object. It rejects inconsistent matches, not all printed replicas.")

    _, draw = deck.slide("Evaluation setup: review labels before running")
    counts = (f"Validation: {summary['ValidationFrames']} frames; "
              f"{summary['ValidationTargetPresent']} present / {summary['ValidationTargetAbsent']} absent.\n"
              f"Evaluation: {summary['EvaluationFrames']} frames; "
              f"{summary['TargetPresent']} present / {summary['TargetAbsent']} absent.") if summary else \
        "Available RGB frames: 126 in eval + 93 in test.\nDefault label manifest: 40 sampled per sequence.\nReviewed/evaluated counts: PENDING."
    deck.card(draw, 65, 205, 710, 555, "Reference labels", counts + "\n"
              "Human presence decision + whole-object box, including hose.\n"
              "Printed replicas are absent, unless the real target is also visible.\n"
              "Unknown labels block evaluation.", 27)
    deck.card(draw, 815, 205, 720, 555, "Held-out protocol",
              "Alternating 20-frame temporal blocks; 2-frame guards on both boundary sides.\n"
              "Sample evenly within retained frames; max 40 per sequence.\n"
              "Both splits must contain positives and negatives.\n"
              "Folder names are not the split. Adjacent views remain correlated; no independent-scene claim.", 27)
    deck.note("This is a custom within-sequence split, not an untouched official test-folder evaluation. "
              "Freeze the manifest before looking at predictions. More independent recordings would improve validity.")

    _, draw = deck.slide("Metrics: correct presence AND correct location")
    deck.card(draw, 65, 205, 710, 555, "Assignment criterion",
              "Correct localization: predicted BOX CENTER lies inside the reference box (inclusive edges).\n"
              "TP: real target + accepted + localized.\n"
              "FP: any accepted region that is not a TP.\n"
              "FN: real target without a localized acceptance.\n"
              "A mislocalized acceptance counts as FP AND FN.", 27)
    deck.card(draw, 815, 205, 720, 555, "Reported measures",
              "Precision = TP / all acceptances; target >= 75%.\n"
              "Recall = TP / real-target frames; target >= 60%.\n"
              "Localization success = localized acceptances / all acceptances; show count and percent.\n"
              "This localization definition equals precision here.\n"
              "Undefined ratios use 0 in tables plus an explicit defined flag.", 27)
    deck.note("TP+FP+FN+TN need not equal frame count because a bad acceptance in a positive frame "
              "is both FP and FN. If asked about IoU, explain that the assignment uses a center criterion instead.")

    _, draw = deck.slide("Results: identical held-out frames and metrics")
    for x, name in [(65, "Baseline"), (815, "Final")]:
        text = metric_text(summary[name]) if summary else \
            "Precision: PENDING\nRecall: PENDING\nTP / FP / FN: PENDING\nLocalization count / percent: PENDING\nRuntime: PENDING\nNo performance claim is justified yet."
        deck.card(draw, x, 220, 710, 480, name, text, 29)
    if summary:
        wrapped(draw, f"Proposal-center recall: {percent(summary['CandidateCenterRecall'])} "
                "(proposal-stage diagnostic, not final detector recall).",
                (80, 740), 1420, 26, MUTED)
    deck.note("Compare absolute counts before percentages. State honestly whether both targets were met. "
              "CPU timings include shared proposal generation, exclude one-time anchor setup and I/O.")

    evidence(deck, root, labels, results, [("TP", "Recognition + correct localization"),
                                          ("replica", "Correct printed-replica rejection")],
             "Visual evidence: success and distractor rejection")
    evidence(deck, root, labels, results, [("FP", "False positive: explain matching evidence"),
                                          ("FN", "False negative: missed or mislocalized target")],
             "Visual evidence: failures")

    image, draw = deck.slide("Threshold trade-off and anchor contribution")
    if not threshold_plot(image, draw, results):
        deck.card(draw, 65, 230, 900, 480, "Validation curve pending",
                  "Threshold curves are generated from validation frames only.\n"
                  "Choose best F1 among thresholds meeting BOTH targets.\n"
                  "If none qualify, select best F1 and explicitly report failure.\n"
                  "Never choose a threshold from evaluation results.", 30)
    wins = str(summary["Final"]["AnchorWins"]) if summary else "PENDING"
    deck.card(draw, 1010, 230, 525, 510, "Three anchors",
              f"Valid top-ranked wins [1,2,3]:\n{wins}\n"
              "Different viewpoints expose different text/edges.\n"
              "Wins measure usage, not causal contribution. A leave-one-anchor-out ablation is needed for that claim.", 27)
    deck.note("Higher thresholds reject more candidates but precision need not be monotonic. "
              "Each anchor's candidates are logged; differences in viewpoint and texture affect usage.")

    _, draw = deck.slide("Reflection: robustness is not physical authenticity")
    deck.card(draw, 65, 205, 710, 555, "Expected sensitivities",
              "Viewpoint: SIFT helps, but a cylinder is not planar.\n"
              "Distance/scale: small labels lose texture; proposal pixels disappear.\n"
              "Illumination: red HSV gates can fail in shadows.\n"
              "Occlusion: fewer inliers, less anchor coverage.\n"
              "Loose anchor crops can match the room, not the object.", 27)
    deck.card(draw, 815, 205, 720, 555, "Limits and next experiments",
              "A printed replica can reproduce excellent SIFT correspondences and a valid homography.\n"
              "RGB matching does NOT prove an object is physical.\n"
              "Red proposals can miss objects or favor red distractors.\n"
              "More proposals / SIFT detail cost processing time.\n"
              "Future: synchronized depth or temporal parallax, independently evaluated.", 27)
    deck.note("Depth is provided but deliberately not implemented here: alignment and timing require validation. "
              "Do not claim geometry alone reliably rejects prints. Discuss observed failure cases after running.")

    _, draw = deck.slide("Contributions, reproduction and demonstration")
    deck.card(draw, 65, 205, 710, 555, "Replace placeholders",
              "[Member 1]: proposals + baseline; presents pipeline.\n"
              "[Member 2]: SIFT + geometry; presents matching.\n"
              "[Member 3]: annotations + evaluation; presents results.\n"
              "Adapt to actual team size and REAL contributions.\n"
              "Every team member presents and explains individual decisions.", 27)
    deck.card(draw, 815, 205, 720, 555, "Reproduce",
              "MATLAB: Project (interactive annotation, tests, experiment and demo).\n"
              "Python: build_report.py generates this PPTX/PDF and ZIP.\n"
              "Results include CSV predictions, candidate ranks, threshold sweeps and label snapshots.\n"
              "Demo: demonstration.mp4 with reference/baseline/final overlays.", 27)
    deck.note("Show the recorded demo. Explain one acceptance and one rejection using scores and threshold. "
              "Before submission replace all member placeholders and confirm measured result counts.")

    _, draw = deck.slide("References", "Outside the 12 main slides")
    references = [
        "D. G. Lowe (2004). Distinctive Image Features from Scale-Invariant Keypoints. IJCV 60, 91-110.",
        "M. A. Fischler & R. C. Bolles (1981). Random Sample Consensus. Communications of the ACM 24(6), 381-395.",
        "MathWorks: detectSIFTFeatures, extractFeatures, matchFeatures, estimateGeometricTransform2D.",
        "Course Project 3 brief and instructor-provided RGB/depth dataset and three anchors.",
        "Implementation assistance: AI assistant using Copilot SDK in VS Code. Review all code and follow course AI-disclosure rules.",
    ]
    y = 220
    for reference in references:
        y = wrapped(draw, reference, (85, y), 1400, 29) + 25
    return deck


def write_outputs(deck: Deck, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    prefix = "project3_draft" if deck.draft else "project3"
    presentation = Presentation()
    presentation.slide_width = Inches(16)
    presentation.slide_height = Inches(9)
    pdf = Canvas(str(output / f"{prefix}.pdf"), pagesize=(960, 540))
    for image, notes in zip(deck.pages, deck.notes):
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(stream, 0, 0, width=Inches(16), height=Inches(9))
        slide.notes_slide.notes_text_frame.text = notes
        stream.seek(0)
        pdf.drawImage(ImageReader(stream), 0, 0, width=960, height=540)
        pdf.showPage()
    presentation.save(output / f"{prefix}.pptx")
    pdf.save()
    deck.pages[2].save(output / "architecture.png")
    if deck.draft:
        frames = [ImageOps.contain(deck.pages[i], (960, 540)) for i in [0, 2, 3, 4, 5, 10]]
        frames[0].save(output / "approach_walkthrough.gif", save_all=True,
                       append_images=frames[1:], duration=4000, loop=0)
    print(f"Built {len(deck.pages)-1} main slides + references: {prefix}.pptx / .pdf")


def package(root: Path, results: Path, output: Path, draft: bool, include_data: bool) -> Path:
    name = "project3_draft_submission.zip" if draft else "project3_submission.zip"
    path = output / name
    files = set(root.glob("*.m")) | set(root.glob("*.py"))
    files.update(root.glob("requirements*.txt"))
    files.update(root.glob("README.md"))
    files.update(root.glob(".gitignore"))
    for directory in ["tests", "annotations"]:
        files.update(p for p in (root / directory).rglob("*") if p.is_file() and "__pycache__" not in p.parts)
    if not draft:
        files.update(p for p in results.rglob("*") if p.is_file())
    if include_data:
        files.update((root / "Anchor_Images").glob("*.png"))
        # Keep every RGB frame so annotation sampling remains reproducible.
        files.update((root / "raw").glob("*\\camera_color_image_raw\\*.png"))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for source in sorted(files):
            if source.is_relative_to(root):
                destination = source.relative_to(root)
            else:
                destination = Path("results") / source.relative_to(results)
            archive.write(source, destination)
        for source in sorted(output.iterdir()):
            if source.suffix in {".pptx", ".pdf", ".gif", ".png"}:
                if draft == ("draft" in source.stem) or source.suffix in {".gif", ".png"}:
                    if not draft and source.name == "approach_walkthrough.gif":
                        continue
                    archive.write(source, Path("presentation") / source.name)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", type=Path, default=ROOT / "results")
    parser.add_argument("--output", type=Path, default=ROOT / "submission")
    parser.add_argument("--include-data", action="store_true",
                        help="Include provided anchors and RGB frames for local ZIP delivery.")
    args = parser.parse_args()
    results = args.results.resolve()
    output = args.output.resolve()
    deck = build_deck(ROOT, results)
    write_outputs(deck, output)
    path = package(ROOT, results, output, deck.draft, args.include_data)
    print(f"Archive: {path}")
    if deck.draft:
        print("DRAFT ONLY: no measured results. Review labels, run MATLAB, exportDemo, then rebuild.")


if __name__ == "__main__":
    main()
